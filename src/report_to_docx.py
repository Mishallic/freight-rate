"""Convert reports/report.md to reports/report.docx, with the figures embedded.

Handles only the Markdown the report uses: headings (#, ##), paragraphs, **bold**,
"- " bullet lists, pipe tables and images (![caption](path)). HTML comments such as
<!-- REWRITE --> are skipped.

Run from the repo root:  python src/report_to_docx.py
"""

import re
from pathlib import Path

from docx import Document
from docx.shared import Inches

REPORTS_DIR = Path(__file__).resolve().parent.parent / "reports"


def add_text(paragraph, text):
    """Add text to a paragraph; parts written as **bold** become bold."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        paragraph.add_run(part).bold = i % 2 == 1  # split() puts the bold parts at odd positions


def add_table(document, lines):
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not set("".join(row)) <= set("-: ")]  # drop the |---| line
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            add_text(table.cell(r, c).paragraphs[0], f"**{text}**" if r == 0 else text)


def add_block(document, lines):
    """Add one block of Markdown (the lines between two blank lines)."""
    lines = [line for line in lines if not line.startswith("<!--")]
    if not lines:
        return
    first = lines[0]
    if first.startswith("#"):
        level = len(first) - len(first.lstrip("#"))
        document.add_heading(first.lstrip("# "), level=level - 1)  # "#" is the title (level 0)
    elif first.startswith("|"):
        add_table(document, lines)
    elif first.startswith("!["):
        caption, path = re.match(r"!\[(.*)\]\((.*)\)", first).groups()
        document.add_picture(str(REPORTS_DIR / path), width=Inches(6))
        document.add_paragraph().add_run(caption).italic = True
    elif first.startswith("- "):
        for line in lines:
            add_text(document.add_paragraph(style="List Bullet"), line[2:])
    else:
        add_text(document.add_paragraph(), " ".join(lines))


def main():
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Inches(1)  # 6.5 in of text width; figures are 6 in
    block = []
    for line in (REPORTS_DIR / "report.md").read_text(encoding="utf-8").splitlines() + [""]:
        if line.strip():
            block.append(line.strip())
        else:
            add_block(document, block)
            block = []
    document.save(REPORTS_DIR / "report.docx")
    print("saved", REPORTS_DIR / "report.docx")


if __name__ == "__main__":
    main()
